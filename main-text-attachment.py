import re
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from pydantic import BaseModel
import pypdf
from docx import Document
import openpyxl
from PIL import Image
import pytesseract
import io

app = FastAPI(title="PrimeServe Advanced DLP Engine")

# Regex Patterns for Data Loss Prevention (DLP)
PAN_PATTERN = r'\b\d{4}[- ]?\d{4}[- ]?\d{4}[- ]?\d{4}\b'
PASSWORD_PATTERN = r'(?i)(password|passwd|pin|secret|credential)\s*[:=]\s*\S+'

class MessagePayload(BaseModel):
    message_text: str

# --- হেল্পার ফাংশন: টেক্সট স্ক্যান করার জন্য ---
def inspect_text(text: str):
    """টেক্সটের মধ্যে কোনো সেনসিটিভ ডেটা আছে কিনা তা যাচাই করে"""
    if re.search(PAN_PATTERN, text):
        return True, "Sensitive Financial Data (Credit/Debit Card Number) detected."
    if re.search(PASSWORD_PATTERN, text):
        return True, "Security Risk: System Password, PIN or Credential exposure detected."
    return False, None

# --- হেল্পার ফাংশন: ফাইল থেকে টেক্সট এক্সট্র্যাক্ট করার জন্য ---
def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    extracted_text = ""
    file_extension = filename.split(".")[-1].lower()

    try:
        # ১. PDF ফাইল প্রসেসিং
        if file_extension == "pdf":
            pdf_file = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(pdf_file)
            for page in reader.pages:
                extracted_text += page.extract_text() or ""
        
        # ২. Word (.docx) ফাইল প্রসেসিং
        elif file_extension == "docx":
            doc_file = io.BytesIO(file_bytes)
            doc = Document(doc_file)
            for para in doc.paragraphs:
                extracted_text += para.text + "\n"

        # ৩. Excel (.xlsx) ফাইল প্রসেসিং
        elif file_extension == "xlsx":
            excel_file = io.BytesIO(file_bytes)
            wb = openpyxl.load_workbook(excel_file, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    extracted_text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"

        # ৪. ইমেজ ফাইল (PNG, JPG) প্রসেসিং via OCR
        elif file_extension in ["png", "jpg", "jpeg"]:
            image = Image.open(io.BytesIO(file_bytes))
            extracted_text = pytesseract.image_to_string(image)
            
        # ৫. প্লেইন টেক্সট ফাইল (.txt)
        elif file_extension == "txt":
            extracted_text = file_bytes.decode("utf-8", errors="ignore")

    except Exception as e:
        # ফাইল রিড করতে সমস্যা হলে লগ বা হ্যান্ডেল করুন
        print(f"Error parsing file {filename}: {str(e)}")
        
    return extracted_text

# --- এন্ডপয়েন্ট ১: শুধু মেসেজ বা টেক্সট স্ক্যান করার জন্য ---
@app.post("/api/v1/scan-message")
def scan_message(payload: MessagePayload):
    is_blocked, reason = inspect_text(payload.message_text)
    return {
        "status": "success",
        "is_blocked": is_blocked,
        "flagged_reason": reason,
        "processed_text": payload.message_text
    }

# --- এন্ডপয়েন্ট ২: মেসেজ এবং অ্যাটাচমেন্ট একসাথে স্ক্যান করার জন্য ---
@app.post("/api/v1/scan-outgoing")
async def scan_outgoing(
    message_text: str = Form(default=""), 
    attachment: UploadFile = File(default=None)
):
    is_blocked = False
    flagged_reason = None
    scanned_sources = []

    # ১. প্রথমে টেক্সট মেসেজ স্ক্যান করুন
    if message_text:
        scanned_sources.append("message_text")
        is_blocked, flagged_reason = inspect_text(message_text)

    # ২. মেসেজ ক্লিন থাকলে এবার অ্যাটাচমেন্ট স্ক্যান করুন
    if not is_blocked and attachment:
        scanned_sources.append(f"attachment ({attachment.filename})")
        file_bytes = await attachment.read()
        
        # ফাইল থেকে টেক্সট বের করুন
        file_text = extract_text_from_file(file_bytes, attachment.filename)
        
        # এক্সট্র্যাক্ট করা টেক্সট স্ক্যান করুন
        if file_text.strip():
            is_blocked, flagged_reason = inspect_text(file_text)

    return {
        "status": "success",
        "is_blocked": is_blocked,
        "flagged_reason": flagged_reason,
        "scanned_sources": scanned_sources
    }

if __name__ == "__main__":
    import uvicorn
    # লোকালহোস্ট হিসেবে সাধারণত 127.0.0.1 স্ট্যান্ডার্ড, তবে আপনার ইচ্ছেমতো পরিবর্তন করতে পারেন
    uvicorn.run(app, host="127.0.0.1", port=5000)